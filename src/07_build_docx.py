"""Fills the course template reports/figs/ANN_Project_Report_Format.docx with this project's
report. Run after 05_final_report.py. Word then refreshes the TOC and exports the PDF
(see README)."""
import csv
import re

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from matplotlib.patches import FancyBboxPatch
from sklearn.metrics import classification_report

from common import ROOT, RUNS, load_json

REP = ROOT / "reports"
FIGS = REP / "figs"
TEMPLATE = FIGS / "ANN_Project_Report_Format.docx"
OUT = REP / "Project29_T_Manoj_ANN_Project_Report.docx"

TITLE = "Paddy Disease Representation using a Convolutional Autoencoder"
STUDENT, ROLL = "T Manoj", "2023004133"
FACULTY = "Anantha Babu"
REPO = None  # final GitHub URL still to come; a highlighted placeholder is used until then
KAGGLE = "https://www.kaggle.com/competitions/paddy-disease-classification"
AE, HEAD = "cae_z128_b32_lr0.1_m0.0", "cae_z128_b32_lr0.1_m0.0__head_b64_lr0.1"
CENTER, LEFT = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT


# ---------------------------------------------------------------- figures made for this report
def box(ax, x, y, w, h, title, body, color, fs=(11, 9.3)):
    ax.add_patch(FancyBboxPatch((x - w / 2, y - h / 2), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
                                fc=color, ec="#333333", lw=0.9))
    ax.text(x, y + h / 2 - 0.12, title, ha="center", va="top", fontsize=fs[0], weight="bold")
    ax.text(x, y - 0.12, body, ha="center", va="center", fontsize=fs[1], linespacing=1.3)


def arrow(ax, a, b):
    ax.annotate("", xy=b, xytext=a, arrowprops=dict(arrowstyle="-|>", color="#333333", lw=1.1))


def draw_workflow(path):
    steps = [("Problem", "Field diagnosis of\npaddy diseases"),
             ("Dataset", "Paddy Doctor (Kaggle)\n10,407 labelled images"),
             ("Preprocessing", "dHash dedup, 70/15/15\n128 x 128, [0, 1], augment"),
             ("Stage 1: CAE", "MSE reconstruction\nmini-batch GD, no labels"),
             ("Stage 2: classifier", "frozen encoder, 128-64-10\nweighted CE, mini-batch GD"),
             ("Experiments", "B1 baseline, E1 main,\nE2 latent, E3 LR / batch"),
             ("Evaluation", "test macro-F1, accuracy,\nAUC, MSE, t-SNE"),
             ("Discussion", "observations and\nlimitations")]
    fig, ax = plt.subplots(figsize=(10, 4.0))
    w, h, dx = 2.75, 1.3, 3.15
    pos = [(i * dx, 2.0) for i in range(4)] + [((3 - i) * dx, 0.0) for i in range(4)]
    colors = ["#eeeeee", "#fce5cd", "#fff2cc", "#cfe2f3", "#d9ead3", "#ead1dc", "#d0e0e3", "#eeeeee"]
    for (x, y), (t, b), c in zip(pos, steps, colors):
        box(ax, x, y, w, h, t, b, c)
    for i in range(7):
        (x0, y0), (x1, y1) = pos[i], pos[i + 1]
        if y0 == y1:
            s = 1 if x1 > x0 else -1
            arrow(ax, (x0 + s * w / 2, y0), (x1 - s * w / 2, y1))
        else:
            arrow(ax, (x0, y0 - h / 2), (x1, y1 + h / 2))
    ax.set_xlim(-1.5, 3 * dx + 1.5); ax.set_ylim(-0.9, 2.9); ax.axis("off")
    fig.tight_layout(); fig.savefig(path, dpi=200); plt.close(fig)


def draw_architecture(path):
    """U-shaped layout: encoder and head on top, decoder running back underneath."""
    fig, ax = plt.subplots(figsize=(10, 5.0))
    enc, lat, dec, head, io = "#cfe2f3", "#fce5cd", "#e6e6e6", "#d9ead3", "#f3f3f3"
    w, dx, top, bot = 1.35, 1.6, 3.5, 0.6
    fs = (10.5, 9)
    row1 = [(0, 1.9, "Input", "128x128x3", io), (1, 1.7, "Conv 3x3/2", "32 ch, ReLU\n64x64x32", enc),
            (2, 1.55, "Conv 3x3/2", "64 ch, ReLU\n32x32x64", enc),
            (3, 1.4, "Conv 3x3/2", "128 ch, ReLU\n16x16x128", enc),
            (4, 1.4, "Bottleneck", "flatten, dense\n32,768 -> 128", enc), (5, 1.25, "Latent z", "128-d", lat),
            (6, 1.25, "Dense 64", "ReLU", head), (7, 1.25, "Dense 10", "Softmax\n10 classes", head)]
    row2 = [(4, 1.4, "Dense", "+ reshape\n16x16x128", dec), (3, 1.55, "ConvT 3x3/2", "64 ch, ReLU\n32x32x64", dec),
            (2, 1.7, "ConvT 3x3/2", "32 ch, ReLU\n64x64x32", dec),
            (1, 1.9, "ConvT 3x3/2", "3 ch, Sigmoid\n128x128x3", io)]
    for row, y in ((row1, top), (row2, bot)):
        for i, h, t, b, c in row:
            box(ax, i * dx, y, w, h, t, b, c, fs)
        for (i0, *_), (i1, *_) in zip(row, row[1:]):
            s = 1 if i1 > i0 else -1
            arrow(ax, (i0 * dx + s * w / 2, y), (i1 * dx - s * w / 2, y))
    arrow(ax, (5 * dx, top - 0.63), (4 * dx + w / 2, bot + 0.3))
    ax.text(0, bot, "Reconstruction\nx̂, MSE vs input", ha="center", va="center", fontsize=9, color="#555555")
    ax.text(2.5 * dx, top + 1.25, "Encoder: trained in Stage 1, frozen in Stage 2", ha="center", fontsize=10.5,
            color="#1f4e79")
    ax.text(6.5 * dx, top + 1.25, "Classifier head: Stage 2", ha="center", fontsize=10.5, color="#38761d")
    ax.text(2.5 * dx, bot - 1.3, "Decoder: Stage 1 only, discarded afterwards", ha="center", fontsize=10.5,
            color="#555555")
    ax.set_xlim(-0.9, 7 * dx + 0.9); ax.set_ylim(bot - 1.5, top + 1.6); ax.axis("off")
    fig.tight_layout(); fig.savefig(path, dpi=200); plt.close(fig)


EQUATIONS = {
    "eq_scale": r"$x' = \frac{x}{255} \in [0, 1], \qquad \tilde{z}_j = \frac{z_j - \mu_j}{\sigma_j}$",
    "eq_mse": r"$L_{MSE} = \frac{1}{N\,C\,H\,W} \sum_{n=1}^{N} \sum_{c,h,w} \left(x_{n,c,h,w} - \hat{x}_{n,c,h,w}\right)^2$",
    "eq_ce": r"$L_{CE} = -\frac{\sum_{n} w_{y_n} \log p_{n,y_n}}{\sum_{n} w_{y_n}}, \qquad"
             r" p_n = \mathrm{softmax}(h(z_n)), \qquad w_k = \frac{N}{K\,N_k}$",
    "eq_sgd": r"$\theta_{t+1} = \theta_t - \eta \cdot \frac{1}{m} \sum_{i \in B_t} \nabla_{\theta} L(x_i;\, \theta_t)$",
    "eq_bp": r"$\delta^{(l)} = \left(W^{(l+1)T} \delta^{(l+1)}\right) \odot g'(s^{(l)}), \qquad"
             r" \frac{\partial L}{\partial W^{(l)}} = \delta^{(l)} \, a^{(l-1)T}$",
}


def render_equation(name, tex):
    path = FIGS / f"_{name}.png"
    fig = plt.figure(figsize=(8, 0.8))
    fig.text(0.5, 0.5, tex, ha="center", va="center", fontsize=15)
    fig.savefig(path, dpi=250, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)
    return path


# ---------------------------------------------------------------- docx helpers
def add_rich(p, text, size=None, italic=False):
    """Adds text with **bold** and *italic* markup."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|\[[A-Z][A-Z ]+\])", text):
        if not tok:
            continue
        marked = len(tok) > 1 and tok.startswith("*") and tok.endswith("*")
        bold = marked and tok.startswith("**")
        r = p.add_run(tok.strip("*") if marked else tok)
        r.bold = True if bold else None
        r.italic = True if italic or (marked and not bold) else None
        if size:
            r.font.size = Pt(size)
    return p


def add_hyperlink(p, url, text=None):
    r_id = p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text or url; t.set(qn("xml:space"), "preserve"); run.append(t)
    link.append(run); p._p.append(link)


def is_heading(el):
    st = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return st is not None and st.get(qn("w:val")).startswith("Heading")


def is_page_break(el):
    return bool(el.xpath('.//w:br[@w:type="page"]'))


def clear_after(el):
    """Removes the template's placeholder paragraphs/tables after a heading."""
    nxt = el.getnext()
    while nxt is not None:
        tag = nxt.tag.split("}")[1]
        if tag == "p" and (is_heading(nxt) or is_page_break(nxt)):
            break
        if tag not in ("p", "tbl"):
            break
        after = nxt.getnext()
        nxt.getparent().remove(nxt)
        nxt = after


class Cursor:
    """Inserts blocks one after another, starting after an anchor element."""

    def __init__(self, doc, anchor):
        self.doc, self.el = doc, anchor

    def _add(self, el):
        self.el.addnext(el)
        self.el = el
        return el

    def para(self, text="", style="Normal", align=None, size=None, italic=False, keep=False):
        p = Paragraph(self._add(OxmlElement("w:p")), self.doc._body)
        p.style = self.doc.styles[style]
        add_rich(p, text, size=size, italic=italic)
        if align is not None:
            p.alignment = align
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    def heading(self, text, level=2):
        p = self.para(style=f"Heading {level}")
        p.add_run(text).font.name = "Times New Roman"
        return p

    def bullets(self, items, numbered=False):
        for it in items:
            self.para(it, style="List Number" if numbered else "List Bullet")

    def image(self, path, width):
        p = self.para(align=CENTER, keep=True)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(path), width=Inches(width))

    def figure(self, path, width, caption):
        self.image(path, width)
        self.para(caption, style="Caption", align=CENTER)

    def equation(self, name):
        self.image(render_equation(name, EQUATIONS[name]), 5.4 if name in ("eq_ce", "eq_bp") else 4.6)

    def table(self, caption, rows, widths, align=None, size=10, bold_last=False):
        self.para(caption, style="Caption", align=CENTER, keep=True)
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._add(t._tbl)
        fix_table_layout(t, widths)
        align = align or [CENTER] * len(rows[0])
        for i, row in enumerate(rows):
            for j, text in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Inches(widths[j])
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.first_line_indent = Inches(0)
                p.alignment = CENTER if i == 0 else align[j]
                bold = i == 0 or (bold_last and i == len(rows) - 1)
                add_rich(p, f"**{text}**" if bold and text and "**" not in text else str(text), size=size)
        # header row repeats across pages
        trpr = t.rows[0]._tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trpr.append(hdr)
        self.para("", size=6)  # breathing room after the table
        return t


def fix_table_layout(t, widths):
    """Fixed column widths (Word ignores cell widths without these) and a little cell padding."""
    tbl_pr = t._tbl.tblPr
    for gc, w in zip(t._tbl.tblGrid.gridCol_lst, widths):
        gc.w = Inches(w)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440))); tbl_w.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", 35), ("left", 80), ("bottom", 35), ("right", 80)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa"); mar.append(e)
    look = tbl_pr.find(qn("w:tblLook"))  # schema order: ... tblLayout, tblCellMar, tblLook
    for el in (layout, mar):
        if look is not None:
            look.addprevious(el)
        else:
            tbl_pr.append(el)


def replace_in_runs(doc, mapping):
    """Plain-text replacements; keys starting with 're:' are regular expressions."""
    for p in doc.paragraphs:
        for r in p.runs:
            for old, new in mapping.items():
                if old.startswith("re:"):
                    if re.search(old[3:], r.text):  # only rewrite matching runs: rewriting turns page breaks into line breaks
                        r.text = re.sub(old[3:], new, r.text)
                elif old in r.text:
                    r.text = r.text.replace(old, new)


def highlight_placeholders(doc):
    """Yellow-highlights [PLACEHOLDER] runs the student still has to fill in."""
    for p in doc.paragraphs:
        for r in p.runs:
            if re.fullmatch(r"\[[A-Z][A-Z ]+\]", r.text.strip()):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def add_toc(cur):
    """Real TOC field restricted to the main body (bookmark _MainBody); Word fills it in."""
    p = cur.para()
    for kind, text in (("begin", None), (None, r'TOC \o "1-2" \h \z \u \b _MainBody'), ("separate", None),
                       (None, None), ("end", None)):
        r = OxmlElement("w:r")
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r.append(fc)
        elif text:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r.append(it)
        else:
            t = OxmlElement("w:t"); t.text = "Right-click and choose Update Field to build the table of contents."
            r.append(t)
        p._p.append(r)


def bookmark_body(doc, first_heading_el):
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), "900"); start.set(qn("w:name"), "_MainBody")
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), "900")
    first_heading_el.insert(1 if first_heading_el.find(qn("w:pPr")) is not None else 0, start)
    doc.paragraphs[-1]._p.append(end)


# ---------------------------------------------------------------- data
def f3(x):
    return f"{x:.3f}"


def load_results():
    from common import load_split
    b1, e1, ae = load_json(RUNS / "B1.json"), load_json(RUNS / f"{HEAD}.json"), load_json(RUNS / f"{AE}.json")
    split = load_split()
    yte, classes = split["y_test"].numpy(), split["classes"]
    pred = np.load(RUNS / f"{HEAD}_test_pred.npy")
    rep = classification_report(yte, pred, target_names=classes, output_dict=True, zero_division=0)
    with open(RUNS / "experiment_log.csv") as f:
        log = {r["run"]: r for r in csv.DictReader(f)}  # last row per run wins
    return b1, e1, ae, rep, classes, log


# ---------------------------------------------------------------- content
def build():
    draw_workflow(FIGS / "workflow.png")
    draw_architecture(FIGS / "architecture.png")
    b1, e1, ae, rep, classes, log = load_results()
    hist = e1["hist"]
    best_ep = int(np.argmin(hist["val"])) + 1
    peak_ep = int(np.argmax(hist["val_acc"])) + 1
    bt, et = b1["test"], e1["test"]
    # differences of the rounded values, so the text matches the tables
    d_f1 = round(et["macro_f1"], 3) - round(bt["macro_f1"], 3)
    d_auc = round(et["macro_auc_ovr"], 3) - round(bt["macro_auc_ovr"], 3)

    doc = Document(str(TEMPLATE))
    headings = {p.text.strip(): p for p in doc.paragraphs if p.style.name.startswith("Heading")}

    # front matter placeholders
    replace_in_runs(doc, {
        "[PROJECT TITLE]": TITLE,
        "A Sample Project Report Template": "Project No. 29  ·  Domain: Agriculture  ·  Unit 2 / Unit 3 (CO2, CO3)",
        "[STUDENT NAME / TEAM MEMBERS]": f"{STUDENT} ({ROLL})",
        "[STUDENT NAME]": STUDENT,
        "[USN / ROLL NUMBER]": ROLL,
        "I/We hereby declare": "I hereby declare",
        "by me/us": "by me",
        "Student Name: ______________________________": f"Student Name: {STUDENT}",
        "USN / Roll No.: ______________________________": f"USN / Roll No.: {ROLL}",
        "[FACULTY NAME]": FACULTY,
        "re:Faculty Guide: _+": f"Faculty Guide: {FACULTY}",
    })

    def section(name, clear=True):
        el = headings[name]._p
        if clear:
            clear_after(el)
        return Cursor(doc, el)

    c = section("ACKNOWLEDGEMENT")
    c.para(f"I thank my faculty guide, {FACULTY}, for guidance throughout this project, and the Department "
           "of Computer Science and Engineering, GITAM University, Bengaluru Campus, for the opportunity to "
           "carry it out as part of the Artificial Neural Networks course.")
    c.para("I acknowledge Petchiammal A., Briskline Kiruba S., D. Murugan and Pandarasamy A., the creators of the "
           "Paddy Doctor dataset, and Kaggle for hosting the Paddy Disease Classification data. The project was "
           "built with the open-source PyTorch, scikit-learn, NumPy and Matplotlib libraries.")
    c.para("An AI coding assistant (Claude, Anthropic) was used to help write the code and draft this report. "
           "All experiments were run on the project code, and every reported number comes from those runs.")

    c = section("ABSTRACT")
    c.para("Diagnosing paddy diseases in the field is slow and depends on expert agronomists. This project "
           "investigates whether a convolutional autoencoder (CAE), trained only to reconstruct images, learns a "
           "compact representation that carries enough disease information to classify paddy images into ten "
           "classes (nine disease and pest-damage conditions plus normal). The 10,407 labelled images of the Kaggle "
           "Paddy Doctor dataset were checked for near-duplicates with a perceptual hash, split 70/15/15 with "
           "stratification and with duplicate clusters kept together, resized to 128 × 128 and scaled to [0, 1]. "
           "A three-layer convolutional encoder compresses each image into a 128-dimensional latent vector, and a "
           "mirrored decoder reconstructs it. The CAE was trained with plain mini-batch gradient descent on the "
           "mean-squared reconstruction error. The frozen encoder's latent vectors were then classified by a small "
           "dense network trained with class-weighted cross-entropy, also by mini-batch gradient descent. "
           "Experiments varied the learning rate, batch size and latent size, with momentum as a separate "
           "ablation. The selected configuration (learning rate 0.1, batch size 32, 150 epochs) reached a test "
           f"reconstruction MSE of 0.0278 and a test macro-F1 of {f3(et['macro_f1'])} (accuracy "
           f"{f3(et['accuracy'])}, macro AUC {f3(et['macro_auc_ovr'])}), compared with {f3(bt['macro_f1'])} for a "
           "PCA plus logistic-regression baseline of the same size. The learned representation therefore helps, "
           "but the reconstructions keep only colour and coarse layout rather than lesions, and the provisional "
           "0.60 macro-F1 target was not reached.")
    c.para("**Keywords:** Artificial Neural Network, Agriculture, Paddy Doctor, Convolutional Autoencoder, "
           "Backpropagation, Mini-Batch Gradient Descent, Macro-F1")

    c = section("TABLE OF CONTENTS")
    add_toc(c)

    c = section("LIST OF FIGURES")
    for f in ["Figure 1. Project workflow", "Figure 2. Dataset class distribution",
              "Figure 3. Convolutional autoencoder and classifier architecture",
              "Figure 4. Classifier head training and validation accuracy",
              "Figure 5. Training and validation loss (autoencoder and classifier head)",
              "Figure 6. Confusion matrix on the test split",
              "Figure 7. Test images and their reconstructions", "Figure 8. t-SNE of test latent vectors"]:
        c.para(f)
    c = section("LIST OF TABLES")
    for t in ["Table 1. Dataset summary", "Table 2. Feature description", "Table 3. Preprocessing steps",
              "Table 4. ANN architecture", "Table 5. Training configuration", "Table 6. Experiment configurations",
              "Table 7. Actual experimental results (test split)", "Table 8. Per-class results on the test split",
              "Table 9. Comparison of all experiments (validation split)", "Table 10. Objective achievement",
              "Table 11. Experiment log"]:
        c.para(t)

    # 1 -------------------------------------------------------------------------------------------
    c = section("1.1 Background")
    c.para("Rice (paddy) is a staple crop, and diseases and pests such as blast, bacterial leaf blight, tungro and "
           "hispa reduce yield when they are not caught early. In the field, diagnosis usually means an expert "
           "inspecting the plants, which is slow and not always available to farmers. Image-based models could "
           "give a first diagnosis from a phone photograph.")
    c.para("This project looks at one part of that question: whether a neural network can learn, without labels, "
           "a compact numerical description of a paddy image that still separates the diseases. If it can, the same "
           "representation could later be reused for search, clustering or classification with few labels.")
    c = section("1.2 Problem Definition")
    c.bullets([
        "**Input:** an RGB field photograph of paddy plants (480 × 640 pixels), resized to 128 × 128 × 3.",
        "**Output / target:** one of 10 labels: bacterial_leaf_blight, bacterial_leaf_streak, "
        "bacterial_panicle_blight, blast, brown_spot, dead_heart, downy_mildew, hispa, normal, tungro.",
        "**Learning task:** (1) unsupervised representation learning: an encoder maps an image x to a latent "
        "vector z of 128 values and a decoder rebuilds x from z; (2) supervised multi-class classification "
        "from z, with the encoder frozen.",
        "**Constraints:** the assigned architecture (convolutional autoencoder, trained from scratch, no "
        "pretrained backbones or transfer learning) and training method (mini-batch gradient descent); only the "
        "labelled Kaggle training images can be scored, because the Kaggle test images are unlabelled; the classes "
        "are imbalanced (5.2 : 1); training on a single laptop GPU (NVIDIA RTX 3050)."])
    c = section("1.3 Objectives")
    c.bullets([
        "To prepare the Paddy Doctor images for training: duplicate check, leakage-free split, resizing and "
        "scaling.",
        "To design and implement a convolutional autoencoder that reconstructs paddy images from a "
        "low-dimensional latent vector.",
        "To train it with mini-batch gradient descent and document batch size, learning rate and convergence.",
        "To classify the 10 classes from the frozen encoder's latent vectors and compare against a simple "
        "baseline of the same size.",
        "To report reconstruction loss, a latent-space visualisation, macro-F1 and per-class results with honest "
        "observations."])
    c = section("1.4 ANN Concepts Used")
    c.bullets([
        "**Convolutional layers:** local receptive fields and weight sharing, with stride 2 for down-sampling.",
        "**Autoencoder:** an encoder–bottleneck–decoder network trained to reproduce its input (unsupervised "
        "learning).",
        "**Transposed convolutions:** learned up-sampling in the decoder.",
        "**Multilayer perceptron (shallow ANN):** the dense classifier head on top of the latent vector.",
        "**Activation functions:** ReLU in hidden layers, sigmoid for the reconstructed pixels, softmax for "
        "class probabilities.",
        "**Backpropagation and mini-batch gradient descent:** gradient computation and the weight update.",
        "**He (Kaiming) normal initialisation** and **early stopping** on validation loss.",
        "**Momentum:** used only in a labelled side comparison (ablation)."])
    c = section("1.5 Project Workflow")
    c.figure(FIGS / "workflow.png", 6.4, "Figure 1. Project workflow")

    # 2 -------------------------------------------------------------------------------------------
    c = section("2.1 Dataset Name and Source")
    p = c.para("**Dataset:** Paddy Doctor: Paddy Disease Classification (Kaggle competition data) [1], [2]")
    p = c.para("**Source:** ")
    add_hyperlink(p, KAGGLE)
    c.para("**Domain:** Agriculture (plant disease diagnosis)")
    c = section("2.2 Dataset Description")
    c.para("The labelled part of the data contains 10,407 RGB JPEG photographs, each 480 × 640 pixels (landscape "
           "or portrait). They are field shots of paddy plants taken with a smartphone, not close-ups of single "
           "leaves. The file train.csv lists, for every image, its label, the paddy variety (10 varieties) and the "
           "plant age (45 to 82 days). A further 3,470 Kaggle test images have no labels and were not used.")
    c.para("The target is the label column, with 10 classes: nine disease or pest-damage conditions and normal. "
           "Only the pixels are used as model input; variety and age are not.")
    c = section("2.3 Dataset Summary")
    c.table("Table 1. Dataset summary", [
        ["Property", "Value", "Source/Method", "Remarks"],
        ["Number of samples", "10,407 labelled images", "train.csv and train_images/",
         "3,470 unlabelled Kaggle test images not used"],
        ["Number of features", "49,152 (128 × 128 × 3 pixels)", "After resizing (raw: 480 × 640 × 3)",
         "Metadata (variety, age) excluded"],
        ["Target variable", "label (disease class)", "train.csv", "One label per image"],
        ["Classes / output range", "10 classes", "EDA (01_eda_split.py)",
         "Imbalanced: 1,764 (normal) to 337 (bacterial panicle blight), 5.2 : 1"],
    ], [1.45, 1.6, 1.65, 1.8])
    c = section("2.4 Feature Description")
    c.table("Table 2. Feature description", [
        ["Feature", "Data type", "Meaning", "Role"],
        ["Image pixels", "uint8 RGB, 480 × 640 × 3", "Field photograph of paddy plants", "Model input"],
        ["label", "Categorical (10 values)", "Disease class assigned by experts", "Target"],
        ["variety", "Categorical (10 values)", "Paddy variety", "Not used (metadata)"],
        ["age", "Integer (45–82)", "Plant age in days", "Not used (metadata)"],
        ["image_id", "String", "Image file name", "Identifier only"],
    ], [1.3, 1.75, 2.1, 1.35])
    c = section("2.5 Dataset Visualization")
    c.para("Figure 2 shows the number of images per class. The four largest classes (normal, blast, hispa, "
           "dead_heart) each have more than 1,400 images, while the three bacterial classes have 337 to 479. This "
           "imbalance is why macro-F1, which weights every class equally, is the headline metric.")
    c.figure(FIGS / "class_distribution.png", 5.4, "Figure 2. Dataset class distribution")

    # 3 -------------------------------------------------------------------------------------------
    c = section("3.1 Data Cleaning")
    c.bullets([
        "**Missing values:** train.csv has no missing values (0 in all four columns), no duplicate image IDs, "
        "and all 10,407 images load correctly.",
        "**Duplicate records:** each image received a 64-bit difference hash (dHash), and images with a Hamming "
        "distance of 2 or less were linked into clusters. 1,203 images (11.6%) belong to a cluster of two or "
        "more, giving 9,697 clusters, and 10 clusters contain mixed labels. Instead of deleting them, each "
        "cluster is kept inside one split (Section 3.5), so a near-copy of a test image is never seen in "
        "training.",
        "**Irrelevant attributes:** variety and age are not used as input, and image_id is an identifier only.",
        "**Invalid values:** none found. Only two image sizes occur (480 × 640 and 640 × 480)."])
    c = section("3.2 Encoding")
    c.para("The text labels were mapped to integers 0–9 in alphabetical order (bacterial_leaf_blight = 0, …, "
           "tungro = 9). The cross-entropy loss takes these integer targets directly, which is equivalent to "
           "one-hot encoding. There are no other categorical inputs.")
    c = section("3.3 Feature Scaling / Normalization")
    c.para("Each image was resized to 128 × 128 with bilinear interpolation (this does not keep the aspect ratio) "
           "and its pixel values were divided by 255, a min-max scaling to [0, 1]. This range matches the decoder's "
           "sigmoid output, so the reconstruction error compares like with like. No per-channel mean/std "
           "normalisation was applied to the images.")
    c.para("For Stage 2, each latent dimension was standardised with the mean μ and standard deviation σ of the "
           "training latents only, because plain gradient descent is sensitive to features on different scales:")
    c.equation("eq_scale")
    c.para("Training batches only were augmented with a random horizontal flip, a rotation within ±15° and a "
           "brightness change within ±20%. Validation and test images were never augmented.")
    c = section("3.4 Outlier Handling")
    c.para("Pixel values are bounded (0 to 255), so no outlier treatment was needed. The data-quality issues that "
           "matter for images, near-duplicates and mixed-label clusters, were handled as described in Section 3.1.")
    c = section("3.5 Train / Validation / Test Split")
    c.para("The split was made first, before anything was fitted, using StratifiedGroupKFold with 20 folds and "
           "random seed 42: 3 folds for test, 3 for validation and 14 for training. The groups are the duplicate "
           "clusters, so no cluster is split. The result is 7,282 training (70.0%), 1,561 validation (15.0%) and "
           "1,564 test (15.0%) images, with the same class proportions in each.")
    c.para("Everything statistical (PCA for the baseline, the latent mean and standard deviation, the class "
           "weights) was fitted on the training split only. Model choices were made on the validation split, and "
           "the test split was used once, for the final evaluation. The split indices are saved to "
           "artifacts/split.npz and reused by every experiment.")
    c = section("3.6 Preprocessing Summary")
    c.table("Table 3. Preprocessing steps", [
        ["Step", "Method Used", "Reason"],
        ["Missing values", "None present (0 nulls); no imputation needed", "Verified in EDA"],
        ["Duplicates", "dHash, Hamming ≤ 2; each cluster kept in one split", "Prevents an inflated test score"],
        ["Categorical encoding", "Label → integer 0–9", "Required by the cross-entropy loss"],
        ["Resizing", "480 × 640 → 128 × 128, bilinear", "Fits GPU memory and training time"],
        ["Scaling", "Pixels ÷ 255 → [0, 1]; latents z-scored with train statistics",
         "Matches the sigmoid output; stable gradient descent"],
        ["Augmentation", "Flip, ±15° rotation, ±20% brightness (train only)", "More variety without leakage"],
        ["Data split", "70 / 15 / 15, stratified, grouped by cluster, seed 42",
         "Same class ratios; no leakage; reproducible"],
    ], [1.5, 2.6, 2.4], align=[LEFT, LEFT, LEFT])

    # 4 -------------------------------------------------------------------------------------------
    c = section("4.1 Architecture Overview")
    c.para("The assigned architecture is a convolutional autoencoder (CAE). Convolutions suit images because each "
           "filter looks at a small neighbourhood and is reused across the whole image. The three encoder "
           "convolutions have only 93,248 weights in total, whereas a single dense layer on the 49,152 input "
           "pixels would need millions. The autoencoder learns without labels: it must squeeze each image through "
           "a 128-value bottleneck and rebuild it, so the bottleneck has to keep the information that matters most "
           "for reconstruction.")
    c.para("The project has two stages. Stage 1 trains the encoder and decoder together on reconstruction. "
           "Stage 2 freezes the encoder, discards the decoder, and trains a small dense classifier head on the "
           "128-value latent vectors. Comparing that classifier with a baseline of the same size shows whether the "
           "unsupervised representation carries disease information.")
    c = section("4.2 Layer Configuration")
    c.table("Table 4. ANN architecture", [
        ["Layer", "Neurons / Shape", "Activation", "Purpose", "Parameters"],
        ["Input", "128 × 128 × 3", "—", "Receive the RGB image", "0"],
        ["Conv 1 (3×3, stride 2, 32 filters)", "64 × 64 × 32", "ReLU", "Low-level edges and colour", "896"],
        ["Conv 2 (3×3, stride 2, 64 filters)", "32 × 32 × 64", "ReLU", "Mid-level texture", "18,496"],
        ["Conv 3 (3×3, stride 2, 128 filters)", "16 × 16 × 128", "ReLU", "Higher-level patterns", "73,856"],
        ["Flatten + Dense (bottleneck)", "128 (latent z)", "Linear", "Compressed representation", "4,194,432"],
        ["Decoder Dense + reshape", "16 × 16 × 128", "ReLU", "Expand the latent vector", "4,227,072"],
        ["ConvTranspose 1 (3×3, stride 2)", "32 × 32 × 64", "ReLU", "Up-sample", "73,792"],
        ["ConvTranspose 2 (3×3, stride 2)", "64 × 64 × 32", "ReLU", "Up-sample", "18,464"],
        ["ConvTranspose 3 (3×3, stride 2)", "128 × 128 × 3", "Sigmoid", "Reconstructed image", "867"],
        ["Head Dense 1", "64", "ReLU", "Classifier hidden layer", "8,256"],
        ["Head Dense 2 (output)", "10", "Softmax", "Class probabilities", "650"],
        ["Total", "Encoder 4,287,680 · Decoder 4,320,195 · Head 8,906", "", "", "8,616,781"],
    ], [1.9, 1.35, 0.85, 1.5, 0.9], align=[LEFT, CENTER, CENTER, LEFT, CENTER], size=9.5, bold_last=True)
    c.para("All convolutions use “same” padding, so each stride-2 layer halves the height and width. Almost all "
           "parameters sit in the two dense layers that connect the 32,768-value feature map to the 128-value "
           "latent vector and back.")
    c = section("4.3 Neural Network Architecture Diagram")
    c.figure(FIGS / "architecture.png", 6.5, "Figure 3. Convolutional autoencoder and classifier architecture")
    c = section("4.4 Forward Data Flow")
    c.para("An image enters as a 128 × 128 × 3 array of values in [0, 1]. Each encoder convolution applies its "
           "filters, adds a bias and passes the result through ReLU, halving the spatial size: 64 × 64 × 32, then "
           "32 × 32 × 64, then 16 × 16 × 128. That map is flattened to 32,768 values and a dense layer projects it "
           "to the 128-value latent vector z.")
    c.para("In Stage 1, the decoder maps z back to 32,768 values, reshapes them to 16 × 16 × 128, and three "
           "transposed convolutions double the size each time back to 128 × 128 × 3. A sigmoid keeps the output "
           "pixels in [0, 1]. In Stage 2, z is standardised and passed through the head: a 64-neuron ReLU layer "
           "and a 10-neuron output layer whose softmax gives the class probabilities. The predicted class is the "
           "one with the highest probability.")
    c = section("4.5 Loss Function")
    c.para("**Stage 1, mean squared error (MSE)** between each image x and its reconstruction x̂, averaged over "
           "the N images in a batch and all C × H × W pixel values:")
    c.equation("eq_mse")
    c.para("MSE is the standard reconstruction loss for real-valued pixels and needs no labels. It does tend to "
           "favour blurry, average-looking outputs, which matters for the results (Section 7.6).")
    c.para("**Stage 2, class-weighted cross-entropy.** Each class k gets a weight w_k inversely proportional to "
           "its training count N_k (K = 10 classes, N training images), so that mistakes on the small bacterial "
           "classes cost as much as mistakes on the large ones:")
    c.equation("eq_ce")
    c = section("4.6 Regularization")
    c.para("No dropout, batch normalisation or L2 weight decay was used, so that the network matches the assigned "
           "plain architecture and optimiser. Overfitting is controlled instead by:")
    c.bullets(["**Early stopping** on validation loss (patience 8), with the best weights restored.",
               "**Data augmentation** of training batches (flip, rotation, brightness).",
               "**The bottleneck itself:** 128 values cannot memorise 49,152-value images."])

    # 5 -------------------------------------------------------------------------------------------
    c = section("5.1 Selected Training Technique")
    c.para("Both stages use **plain mini-batch gradient descent with backpropagation** (PyTorch "
           "torch.optim.SGD with momentum 0 and no weight decay). In each step the parameters θ move against the "
           "loss gradient averaged over a mini-batch B_t of m examples, with learning rate η:")
    c.equation("eq_sgd")
    c.para("The training set is reshuffled every epoch and split into consecutive mini-batches, so each epoch "
           "visits every training image once.")
    c = section("5.2 Training Configuration")
    c.table("Table 5. Training configuration", [
        ["Parameter", "Actual Value"],
        ["Optimizer / training method", "Mini-batch gradient descent (SGD, momentum 0), both stages"],
        ["Learning rate", "Stage 1: 0.1 · Stage 2: 0.1 (both chosen from 0.1 / 0.01 / 0.001)"],
        ["Batch size", "Stage 1: 32 (chosen from 32 / 64 / 128) · Stage 2: 64"],
        ["Epochs", f"Stage 1: 150 (best epoch 150) · Stage 2: {len(hist['val'])} run, best epoch {best_ep} "
                   "(maximum 50)"],
        ["Loss function", "Stage 1: MSE · Stage 2: class-weighted cross-entropy"],
        ["Weight initialization", "He (Kaiming) normal, biases zero"],
        ["Early stopping", "Patience 8 on validation loss, best weights restored (both stages)"],
        ["Latent size", "128 (chosen from 64 / 128 / 256)"],
        ["Random seed", "42 (split, initialisation, shuffling)"],
        ["Hardware / software", "NVIDIA RTX 3050 Laptop GPU; PyTorch 2.14 (CUDA 12.6), scikit-learn"],
        ["Time per epoch", f"Stage 1: {float(log[AE]['sec_per_epoch']):.1f} s (GPU shared with a second run; "
                           "about 5.4 s at batch 64 when alone) · Stage 2: under 1 s"],
    ], [2.0, 4.5], align=[LEFT, LEFT])
    c = section("5.3 Backpropagation / Weight Update")
    c.bullets([
        "**Forward pass:** a mini-batch flows through the network as in Section 4.4, and every layer keeps its "
        "inputs and pre-activations.",
        "**Loss calculation:** the MSE (Stage 1) or weighted cross-entropy (Stage 2) is computed for the batch.",
        "**Gradient computation:** backpropagation applies the chain rule from the output layer back to the "
        "first layer. For layer l with weights W, pre-activation s, activation a and activation function g, the "
        "error signal δ and the weight gradient are:",
    ])
    c.equation("eq_bp")
    c.bullets([
        "**Weight update:** each parameter takes one step against its batch-averaged gradient (Section 5.1). "
        "PyTorch's autograd computes these gradients automatically, including for convolutions, where the "
        "gradient of a shared filter is summed over every position it was applied to.",
        "**Stage 2:** the encoder is frozen, so gradients only reach the 8,906 head parameters."])
    c = section("5.4 Why This Training Method?")
    c.para("Mini-batch gradient descent was assigned, and it also fits this problem. Full-batch gradient descent "
           "would need all 7,282 training images (49,152 values each) and their activations in GPU memory at "
           "once, which a 4 GB laptop GPU cannot hold. Per-image updates (pure stochastic gradient descent) would "
           "waste the GPU's parallelism. Mini-batches of 32 to 128 images fit in memory, give 57 to 228 updates "
           "per epoch, and let each batch be augmented independently.")
    c.para("Using plain gradient descent without momentum or adaptive learning rates makes the effect of the "
           "learning rate and batch size directly visible, which is the point of experiment E3. Its weakness, "
           "slow convergence, was a known risk; momentum is reported only as a separate ablation.")

    # 6 -------------------------------------------------------------------------------------------
    c = section("6.1 Experimental Design")
    c.para("Kept constant across experiments: the saved split, seed 42, the preprocessing, early stopping "
           "(patience 8), and the encoder depth. When encoders were compared, the head was always trained the same "
           "way (learning rate 0.01, batch size 64), so differences come from the encoder.")
    c.para("Changed one factor at a time, judged on the validation split: the learning rate, the batch size and "
           "the latent size of the autoencoder, each for 50 epochs. Because every plain-SGD run was still improving "
           "at 50 epochs, the selected configuration was retrained for 150 epochs, and then the head's learning "
           "rate was tuned. Momentum 0.9 is a labelled side comparison only. The optional fine-tuning experiment "
           "(E4, encoder unfrozen) was not run.")
    c = section("6.2 Experiment Configuration Table")
    c.table("Table 6. Experiment configurations", [
        ["Exp.", "Architecture", "Learning Rate", "Batch Size", "Epochs", "Purpose"],
        ["B1", "PCA (128) on 64 × 64 pixels + logistic regression", "n/a (L-BFGS)", "Full batch",
         "≤ 2,000 iterations", "Baseline"],
        ["E1", "CAE, z = 128, + head 128–64–10", "CAE 0.1 · head 0.1", "CAE 32 · head 64",
         f"CAE 150 · head {len(hist['val'])}", "Main model"],
        ["E2", "CAE, z = 64 and z = 256", "0.1", "64", "50", "Test architecture (latent size)"],
        ["E3a", "CAE, z = 128", "0.1 / 0.01 / 0.001", "64", "50", "Test learning rate"],
        ["E3b", "CAE, z = 128", "0.1", "32 / 64 / 128", "50", "Test batch size"],
        ["M", "CAE, z = 128, momentum 0.9", "0.1 / 0.01", "64", "50 (and 150 at 0.1)", "Momentum ablation"],
        ["H", "Head on the E1 encoder", "0.1 / 0.01 / 0.001", "64", "≤ 50 (early stop)", "Head learning rate"],
    ], [0.5, 1.75, 1.05, 0.95, 1.05, 1.2], size=9.5)
    # 6.3 keeps the template's procedure list

    # 7 -------------------------------------------------------------------------------------------
    c = section("7.1 Overall Performance")
    c.para("Table 7 gives the final results. Precision, recall and F1 are macro-averaged over the 10 classes, and "
           "AUC is the macro one-vs-rest ROC AUC. Only B1 and E1 were evaluated on the test split, once, so that "
           "the test set stayed untouched during tuning. The other experiments are compared on the validation "
           "split in Section 7.5.")
    c.table("Table 7. Actual experimental results (test split)", [
        ["Experiment", "Accuracy", "Precision", "Recall", "F1", "AUC", "Notes"],
        ["B1 (test)", f3(bt["accuracy"]), f3(bt["macro_precision"]), f3(bt["macro_recall"]),
         f3(bt["macro_f1"]), f3(bt["macro_auc_ovr"]), "PCA-128 + logistic regression baseline"],
        ["E1 (test)", f3(et["accuracy"]), f3(et["macro_precision"]), f3(et["macro_recall"]),
         f3(et["macro_f1"]), f3(et["macro_auc_ovr"]),
         f"+{d_f1:.3f} macro-F1 over B1; below the 0.60 target"],
        ["B1 (validation)", f3(b1["val"]["accuracy"]), f3(b1["val"]["macro_precision"]),
         f3(b1["val"]["macro_recall"]), f3(b1["val"]["macro_f1"]), "—", "Reference"],
        ["E1 (validation)", f3(e1["val"]["accuracy"]), f3(e1["val"]["macro_precision"]),
         f3(e1["val"]["macro_recall"]), f3(e1["val"]["macro_f1"]), "—", "Optimistic: used for model selection"],
    ], [1.15, 0.75, 0.75, 0.7, 0.6, 0.6, 1.95], align=[LEFT] + [CENTER] * 5 + [LEFT], size=9.5)
    c.para(f"E1 beats the baseline on every metric, by {d_f1:.3f} macro-F1 and {d_auc:.3f} AUC, at the same "
           "representation size of 128 values. "
           "The PRD's acceptance criterion E1 > B1 is met; the provisional target of 0.60 macro-F1 is not. The "
           "test reconstruction MSE of the autoencoder is 0.0278.")
    c = section("7.2 Training and Validation Accuracy")
    c.para("The autoencoder has no accuracy, because Stage 1 has no labels. Figure 4 shows the classifier head's "
           f"accuracy per epoch. Training accuracy (measured on the training batches) climbs steadily to "
           f"{hist['train_acc'][-1]:.3f}, while validation accuracy rises more slowly, reaching "
           f"{hist['val_acc'][best_ep - 1]:.3f} at the best validation-loss epoch ({best_ep}) and peaking at "
           f"{max(hist['val_acc']):.3f} at epoch {peak_ep}.")
    c.figure(FIGS / "accuracy_curves.png", 5.6, "Figure 4. Classifier head training and validation accuracy")
    c = section("7.3 Training and Validation Loss")
    c.figure(FIGS / "loss_curves.png", 6.4,
             "Figure 5. Training and validation loss (left: autoencoder MSE; right: classifier head)")
    c.bullets([
        "**Autoencoder plateau:** with plain gradient descent the MSE drops quickly to about 0.043 and then stays "
        "flat for about 25 epochs. That level matches a network that outputs roughly an average green image.",
        "**Break-out:** after about epoch 25 the loss falls steadily again. This plateau explains why the 50-epoch "
        "runs all looked similar, and why training longer mattered more than any single hyperparameter.",
        "**Generalisation:** final MSE is 0.0254 on training, 0.0279 on validation and 0.0278 on test, so the "
        "validation curve tracks the training curve closely. Small validation spikes after about epoch 90 "
        "suggest a learning rate of 0.1 is close to the stable limit.",
        f"**Classifier head:** training loss keeps falling, but validation loss bottoms out at epoch {best_ep} "
        "and then rises, so the head starts to overfit and early stopping restores the epoch-"
        f"{best_ep} weights."])
    c = section("7.4 Confusion Matrix / Regression Metrics")
    c.figure(FIGS / "confusion_matrix.png", 4.9, "Figure 6. Confusion matrix on the test split")
    rows = [["Class", "Precision", "Recall", "F1", "Support"]]
    for k in classes:
        r = rep[k]
        rows.append([k, f3(r["precision"]), f3(r["recall"]), f3(r["f1-score"]), f"{int(r['support']):,}"])
    m = rep["macro avg"]
    rows.append(["Macro average", f3(m["precision"]), f3(m["recall"]), f3(m["f1-score"]), f"{int(m['support']):,}"])
    c.table("Table 8. Per-class results on the test split", rows, [2.2, 1.0, 1.0, 1.0, 1.0],
            align=[LEFT] + [CENTER] * 4, bold_last=True)
    c.bullets([
        "**bacterial_leaf_blight (F1 0.351)** is the weakest class. Recall is acceptable (0.472), but precision is "
        "only 0.279: many blast, dead_heart, hispa and normal images are predicted as this class. The class weights "
        "push predictions toward small classes, and its yellowing leaf tips resemble other yellow-tinted images at "
        "this resolution.",
        "**bacterial_panicle_blight (F1 0.433)** is the smallest class (337 images). The disease shows on the "
        "panicle, a small region that 128 × 128 resizing and a blurry latent cannot resolve.",
        "**normal, hispa and blast**, the largest classes, are often confused with each other: 31 normal test images "
        "were predicted as hispa and 28 blast images as normal. Their differences are small lesions or scrape "
        "marks, which is exactly the detail the reconstructions lose.",
        "Small classes have higher recall than precision, the expected effect of the class-weighted loss."])
    c = section("7.5 Experiment Comparison")
    c.para("Table 9 compares every experiment on the validation split. For the autoencoder runs, “head val "
           "macro-F1” is a head trained the same way (learning rate 0.01, batch 64) on that run's frozen latents.")

    def v(run, key):
        r = log.get(run, {})
        return f"{float(r[key]):.4f}" if r.get(key) else "–"

    def hf(ae_run):
        return v(f"{ae_run}__head_b64_lr0.01", "val_macro_f1")

    c.table("Table 9. Comparison of all experiments (validation split)", [
        ["Exp.", "Change (others at default)", "Epochs", "Best val MSE", "Head val macro-F1", "Observation"],
        ["Start", "η 0.01, batch 64, z 128", "50", v("cae_z128_b64_lr0.01_m0.0", "best_val_loss"),
         hf("cae_z128_b64_lr0.01_m0.0"), "Still decreasing at epoch 50"],
        ["E3a", "η 0.1", "50", v("cae_z128_b64_lr0.1_m0.0", "best_val_loss"), hf("cae_z128_b64_lr0.1_m0.0"),
         "Faster; stable"],
        ["E3a", "η 0.001", "50", v("cae_z128_b64_lr0.001_m0.0", "best_val_loss"), "–", "Far too slow"],
        ["E3b", "η 0.1, batch 32", "50", "0.0386", "0.3061", "Best plain-SGD run at 50 epochs"],
        ["E3b", "η 0.1, batch 128", "50", v("cae_z128_b128_lr0.1_m0.0", "best_val_loss"), "–",
         "Fewer updates per epoch"],
        ["E2", "η 0.1, z 64", "50", v("cae_z64_b64_lr0.1_m0.0", "best_val_loss"), hf("cae_z64_b64_lr0.1_m0.0"), ""],
        ["E2", "η 0.1, z 256", "50", v("cae_z256_b64_lr0.1_m0.0", "best_val_loss"), hf("cae_z256_b64_lr0.1_m0.0"),
         "Best latent size for classification"],
        ["M", "η 0.1, momentum 0.9", "50", "0.0254", "0.4411", "Much faster convergence"],
        ["M", "η 0.01, momentum 0.9", "50", v("cae_z128_b64_lr0.01_m0.9", "best_val_loss"), "–", ""],
        ["E1 (AE)", "η 0.1, batch 32, z 128", "150", v(AE, "best_val_loss"), hf(AE), "Selected encoder"],
        ["M", "η 0.1, momentum 0.9", "150", v("cae_z128_b64_lr0.1_m0.9", "best_val_loss"),
         hf("cae_z128_b64_lr0.1_m0.9"), "Lower MSE, worse head"],
        ["H", "Head η 0.001 on E1 encoder", "50", "–", v(f"{AE}__head_b64_lr0.001", "val_macro_f1"), "Too slow"],
        ["H", "Head η 0.1 on E1 encoder", str(len(hist["val"])), "–", f"{e1['val']['macro_f1']:.4f}",
         "Selected head (final E1)"],
    ], [0.6, 1.6, 0.6, 0.85, 0.95, 1.9], align=[CENTER, LEFT, CENTER, CENTER, CENTER, LEFT], size=9)
    c.para("Two 50-epoch values (batch 32 and momentum 0.9 at η 0.1) are taken from the experiment log rows "
           "written before the 150-epoch runs reused those run names.", size=10, italic=True)
    c.bullets([
        "**Learning rate** had the largest effect: η = 0.001 barely moved in 50 epochs, and η = 0.1 was fastest "
        "while staying stable.",
        "**Smaller batches** converged faster at a fixed learning rate, because batch 32 gives twice as many "
        "updates per epoch as batch 64.",
        "**Latent size:** at 50 epochs, a larger latent gave a better classifier (0.207, 0.233, 0.299 for 64, 128, "
        "256), even though the reconstruction error was about the same.",
        "**Training length** mattered most: the selected plain-SGD encoder improved from 0.306 to 0.465 head "
        "macro-F1 between 50 and 150 epochs.",
        "**Momentum** reached a much lower MSE (0.0223) but a worse classifier after 150 epochs (0.419 vs 0.465), "
        "so lower reconstruction error did not mean better disease features."])
    c.heading("7.6 Reconstruction Quality and Latent Space")
    c.figure(FIGS / "reconstructions.png", 5.6, "Figure 7. Test images (rows 1 and 3) and their reconstructions "
                                                "directly below (rows 2 and 4)")
    c.para("The reconstructions keep only the overall colour and the coarse light/dark layout. Individual leaves, "
           "leaf edges and lesions are **not** preserved, and a regular grid pattern is visible. The grid likely "
           "comes from the dense layer that expands z back to a 16 × 16 map, followed by stride-2 transposed "
           "convolutions, which are known to produce such artifacts [7]. **The PRD criterion “leaf shape and "
           "lesion regions visibly preserved” is not met.**")
    c.figure(FIGS / "tsne.png", 5.2, "Figure 8. t-SNE of test latent vectors, coloured by class")
    c.para("The t-SNE projection [6] of the test latents is mostly one mixed cloud, so the classes are not cleanly "
           "separated. There are weak trends (normal, hispa and bacterial_leaf_streak toward the upper left) and a "
           "few small tight clusters, mainly of blast and hispa. These are probably photos from the same plot or "
           "session that look nearly identical without passing the duplicate threshold, and they may make some "
           "classes look easier than they are.")

    # 8 -------------------------------------------------------------------------------------------
    c = section("8.1 Interpretation of Results")
    c.para(f"The learned representation helps, but modestly. E1 reaches {f3(et['macro_f1'])} test macro-F1 against "
           f"{f3(bt['macro_f1'])} for B1, and {f3(et['macro_auc_ovr'])} AUC against {f3(bt['macro_auc_ovr'])}, "
           "with the same 128-value representation. The AUC of 0.847 shows the classifier ranks the correct class "
           "fairly well even when its top choice is wrong.")
    c.para("The reconstructions and t-SNE explain why the gain is limited. The latent vector mostly encodes "
           "colour, brightness and coarse texture rather than lesion shapes, which is the PRD's listed risk "
           "“latent captures colour/background, not lesions”. Those global statistics are still more "
           "discriminative than linear PCA components, which is where the improvement over B1 comes from.")
    c = section("8.2 Effect of Architecture")
    c.para("Only the latent size was varied. At 50 epochs a larger latent gave a better classifier (0.207, 0.233 "
           "and 0.299 head macro-F1 for 64, 128 and 256), while the reconstruction error barely changed. A wider "
           "bottleneck keeps more information that the reconstruction loss does not reward but the classifier can "
           "use.")
    c.para("The bigger architectural limit is the bottleneck design. One dense layer squeezes a 32,768-value "
           "feature map into 128 numbers after only three convolutions, and it holds almost all of the 8.6 million "
           "parameters. Fine spatial detail is lost there, as the reconstructions show.")
    c = section("8.3 Effect of Optimization / Training")
    c.para("The learning rate was the most important setting, followed by the training length. Plain gradient "
           "descent sits on a plateau near the “average image” solution for about 25 epochs before it starts to "
           "learn structure, so short runs undersell it. Smaller batches helped at a fixed learning rate because "
           "they give more updates per epoch. Momentum 0.9 escaped the plateau much faster and reached a lower "
           "reconstruction error, but its latents classified slightly worse after 150 epochs. The main model "
           "therefore uses plain gradient descent, as assigned, and momentum is reported only as an ablation.")
    c = section("8.4 Overfitting / Underfitting")
    c.bullets([
        "**Autoencoder: underfitting rather than overfitting.** Training and validation MSE stay close (0.0254 "
        "vs 0.0279) and were still falling at epoch 150, and test MSE (0.0278) matches validation. More "
        "training or a larger model could still improve it.",
        f"**Classifier head: overfits quickly.** Validation loss is lowest at epoch {best_ep} while training "
        f"accuracy keeps rising to {hist['train_acc'][-1]:.3f}. Validation accuracy kept improving slightly after "
        f"that (peak {max(hist['val_acc']):.3f} at epoch {peak_ep}), so stopping on validation loss, as the PRD "
        "specified, may have stopped a little early for accuracy.",
        f"**Validation vs test:** E1 scored {f3(e1['val']['macro_f1'])} macro-F1 on validation and "
        f"{f3(et['macro_f1'])} on test. Many configurations were compared on the same validation split, so the "
        "validation number is optimistic. The test number, evaluated once, is the one to quote."])
    c = section("8.5 Objective Achievement")
    c.table("Table 10. Objective achievement", [
        ["Objective", "Status", "Evidence"],
        ["Prepare the data", "Achieved", "Dedup, grouped 70/15/15 split, saved indices (Section 3)"],
        ["Build a CAE", "Achieved, with a limitation",
         "Trains and converges, but reconstructions lose lesions (Figure 7)"],
        ["Train with mini-batch GD and document it", "Achieved", "Learning rate, batch size and convergence "
                                                                  "results (Tables 5, 9; Figure 5)"],
        ["Classify from frozen latents vs a baseline", "Achieved",
         f"E1 {f3(et['macro_f1'])} vs B1 {f3(bt['macro_f1'])} test macro-F1 (Table 7)"],
        ["Report loss, latent space, macro-F1, per-class", "Achieved", "Figures 5–8, Tables 7–8"],
        ["Macro-F1 ≥ 0.60 (provisional target)", "Not met", f"{f3(et['macro_f1'])} on test"],
    ], [2.2, 1.4, 2.9], align=[LEFT, CENTER, LEFT], size=9.5)

    # 9 -------------------------------------------------------------------------------------------
    c = section("9. Limitations")
    c.bullets([
        "**Dataset size and representativeness:** all images come from one dataset source, taken in similar "
        "field conditions. Performance on photos from other regions, cameras or seasons is unknown.",
        "**Class imbalance:** the classes range from 1,764 to 337 images. Class weighting improves recall on "
        "the small classes but costs precision (bacterial_leaf_blight precision 0.279).",
        "**Noise and near-duplicates:** 10 duplicate clusters had mixed labels, and the tight t-SNE clusters "
        "suggest some near-identical photos remain, possibly split across training and test.",
        "**Computational limits:** training ran on a laptop GPU, so images were reduced to 128 × 128 without "
        "keeping the aspect ratio. Lesions a few pixels wide in the original are mostly lost before the model "
        "sees them.",
        "**Limited hyperparameter search:** one factor at a time, 50-epoch sweeps while plain gradient descent "
        "was still improving, a single seed (42) with no repeats, and no fine-tuning run (E4).",
        "**Interpretability:** the 128 latent values have no direct meaning, and the blurry reconstructions "
        "show that the latent space does not represent lesions explicitly.",
        "**Generalisation:** the encoder was trained without labels, so nothing told it which details are "
        "diagnostic; MSE rewards average-looking reconstructions rather than disease features."])

    # 10 / 11 / appendices ---------------------------------------------------------------------------
    c = section("10. Conclusion")
    c.para("This project trained a convolutional autoencoder from scratch on 10,407 Paddy Doctor field images "
           "with plain mini-batch gradient descent, then classified ten disease classes from its frozen "
           "128-value latent vectors. After a duplicate-aware 70/15/15 split and one-factor-at-a-time tuning, the "
           "selected model (learning rate 0.1, batch size 32, 150 epochs) reached a test reconstruction MSE of "
           f"0.0278 and a test macro-F1 of {f3(et['macro_f1'])} (accuracy {f3(et['accuracy'])}, AUC "
           f"{f3(et['macro_auc_ovr'])}), beating a PCA plus logistic-regression baseline of the same size "
           f"({f3(bt['macro_f1'])} macro-F1).")
    c.para("The unsupervised representation therefore carries useful disease signal, mostly colour and coarse "
           "texture. It does not preserve lesions, the reconstruction criterion was not met, and the provisional "
           "0.60 macro-F1 target was not reached. Learning rate and training length were the decisive factors for "
           "plain mini-batch gradient descent.")
    c = section("10.1 Future Work")
    c.bullets([
        "Train a 150-epoch model with a 256-value latent, the best latent size for classification at 50 epochs.",
        "Run E4: unfreeze the encoder with a low learning rate to find the upper bound for this architecture.",
        "Add a fourth convolution layer or use a convolutional bottleneck, so less spatial detail is discarded.",
        "Use larger inputs (for example 224 × 224) or crop patches around leaves.",
        "Repeat the key runs over several seeds and report the mean and standard deviation.",
        "Test on photos from other regions and devices before any field use."])

    c = section("11. References")
    refs = [
        "[1] Petchiammal A., Briskline Kiruba S., Murugan D., & Pandarasamy A. (2023). Paddy Doctor: A visual "
        "image dataset for automated paddy disease classification and benchmarking. In Proceedings of the 6th "
        "Joint International Conference on Data Science & Management of Data (CODS-COMAD 2023). ACM. "
        "DOI: 10.1145/3570991.3570994",
        f"[2] Kaggle. (2022). Paddy Doctor: Paddy Disease Classification [Competition data]. {KAGGLE}",
        "[3] Rumelhart, D. E., Hinton, G. E., & Williams, R. J. (1986). Learning representations by "
        "back-propagating errors. Nature, 323, 533–536. DOI: 10.1038/323533a0",
        "[4] Hinton, G. E., & Salakhutdinov, R. R. (2006). Reducing the dimensionality of data with neural "
        "networks. Science, 313(5786), 504–507. DOI: 10.1126/science.1127647",
        "[5] He, K., Zhang, X., Ren, S., & Sun, J. (2015). Delving deep into rectifiers: Surpassing human-level "
        "performance on ImageNet classification. In Proceedings of the IEEE International Conference on Computer "
        "Vision (ICCV), 1026–1034.",
        "[6] van der Maaten, L., & Hinton, G. (2008). Visualizing data using t-SNE. Journal of Machine Learning "
        "Research, 9, 2579–2605.",
        "[7] Odena, A., Dumoulin, V., & Olah, C. (2016). Deconvolution and checkerboard artifacts. Distill. "
        "DOI: 10.23915/distill.00003",
        "[8] Goodfellow, I., Bengio, Y., & Courville, A. (2016). Deep Learning. MIT Press.",
        "[9] Paszke, A., et al. (2019). PyTorch: An imperative style, high-performance deep learning library. "
        "Advances in Neural Information Processing Systems, 32.",
        "[10] Pedregosa, F., et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning "
        "Research, 12, 2825–2830."]
    for r in refs:
        p = c.para(r, align=LEFT)
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)

    c = section("APPENDIX A – SOURCE CODE / IMPORTANT CODE SNIPPETS")
    p = c.para("**GitHub repository:** " + ("" if REPO else "[GITHUB REPOSITORY LINK]"), align=LEFT)
    if REPO:
        add_hyperlink(p, REPO)

    c = section("APPENDIX B – ADDITIONAL RESULTS")
    c.para("Table 11 lists every run in runs/experiment_log.csv (the latest row for each run name). Times per "
           "epoch are higher for runs that shared the GPU with a second run.")
    rows = [["Run", "Stage", "Epochs", "Best val loss", "Val macro-F1", "s / epoch"]]
    for name, r in log.items():
        if r["stage"] == "baseline":
            continue
        short = name.replace("cae_", "").replace("__head", " → head")
        rows.append([short, r["stage"], r["epochs_run"],
                     f"{float(r['best_val_loss']):.4f}" if r["best_val_loss"] else "–",
                     f"{float(r['val_macro_f1']):.3f}" if r["val_macro_f1"] else "–",
                     f"{float(r['sec_per_epoch']):.1f}" if r["sec_per_epoch"] else "–"])
    c.table("Table 11. Experiment log", rows, [2.75, 0.6, 0.6, 0.95, 0.9, 0.7],
            align=[LEFT] + [CENTER] * 5, size=8.5)

    highlight_placeholders(doc)
    bookmark_body(doc, headings["1. Problem Statement"]._p)
    doc.core_properties.title = TITLE
    doc.core_properties.author = f"{STUDENT} ({ROLL})"
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
